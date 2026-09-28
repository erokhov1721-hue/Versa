import openpyxl

from versa.core.excel_layout import discover_layout, discover_vat_rate


def test_discovers_seven_participant_blocks(sample_path):
    wb = openpyxl.load_workbook(sample_path, data_only=True)
    ws = wb["3_ ПОДРОБНАЯ"]

    layout = discover_layout(ws)

    assert len(layout.participants) == 7
    names = [p.name for p in layout.participants]
    assert names == [
        'ООО "ПАРАЛЛЕЛЬ"',
        'ООО "Р.И.К. ИНЖИНИРИНГ"',
        'ООО "РУТЕК"',
        'ООО "БЮРО КОНСТРАКШН"',
        'АО "ГЭС КОНСТРАКШН"',
        'АО "ФОДД"',
        'ООО "ЭРБЕК"',
    ]


def test_rik_block_has_no_pct_column_others_do(sample_path):
    wb = openpyxl.load_workbook(sample_path, data_only=True)
    ws = wb["3_ ПОДРОБНАЯ"]

    layout = discover_layout(ws)
    by_name = {p.name: p for p in layout.participants}

    assert by_name['ООО "Р.И.К. ИНЖИНИРИНГ"'].pct_col is None
    assert by_name['ООО "ПАРАЛЛЕЛЬ"'].pct_col is not None
    assert by_name['ООО "РУТЕК"'].pct_col is not None


def test_baseline_block_present_but_columns_located(sample_path):
    wb = openpyxl.load_workbook(sample_path, data_only=True)
    ws = wb["3_ ПОДРОБНАЯ"]

    layout = discover_layout(ws)

    assert layout.baseline.unit_price_cols.total == 13  # column M
    assert layout.baseline.total_cols.total == 17  # column Q


def test_discovers_vat_rate_from_participant_header(sample_path):
    wb = openpyxl.load_workbook(sample_path, data_only=True)
    ws = wb["3_ ПОДРОБНАЯ"]

    assert discover_vat_rate(ws) == 22.0


def test_vat_rate_defaults_when_not_found():
    class FakeCell:
        value = None

    class FakeWs:
        max_row = 13
        max_column = 5

        def cell(self, row, column):
            return FakeCell()

    assert discover_vat_rate(FakeWs()) == 22.0
