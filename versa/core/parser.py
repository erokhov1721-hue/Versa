from __future__ import annotations

from pathlib import Path

import openpyxl

from versa.core.excel_layout import Layout, PriceCols, discover_layout
from versa.core.hierarchy import RawRow, build_section_tree
from versa.core.models import (
    BaselinePrice,
    ComparisonMode,
    ParticipantInfo,
    Position,
    PositionPrice,
    PriceBreakdown,
    PriceStatus,
    Section,
    Tender,
    Warning,
    WarningType,
)
from versa.core.status import classify_status

SHEET_NAME = "3_ ПОДРОБНАЯ"


def _num(value, row: int, col: int, warnings: list[Warning]) -> float | None:
    if value is None:
        return None
    if isinstance(value, str):
        if value.startswith("#"):
            warnings.append(Warning(
                type=WarningType.FORMULA_ERROR, row=row, participant_id=None,
                message=f"Ошибка формулы {value!r} в ячейке ({row}, {col})",
            ))
            return None
        try:
            return float(value.replace(",", "."))
        except ValueError:
            return None
    return float(value)


def _read_breakdown(ws, row: int, cols: PriceCols, warnings: list[Warning]) -> PriceBreakdown:
    return PriceBreakdown(
        materials=_num(ws.cell(row=row, column=cols.materials).value, row, cols.materials, warnings),
        works=_num(ws.cell(row=row, column=cols.works).value, row, cols.works, warnings),
        overhead=_num(ws.cell(row=row, column=cols.overhead).value, row, cols.overhead, warnings),
        total=_num(ws.cell(row=row, column=cols.total).value, row, cols.total, warnings),
    )


def _read_header(ws) -> tuple[str, str, str]:
    subject = str(ws.cell(row=3, column=4).value or "").strip()
    object_name = str(ws.cell(row=4, column=4).value or "").strip()
    address = str(ws.cell(row=5, column=4).value or "").strip()
    return subject, object_name, address


def _read_participants(ws, layout: Layout) -> list[ParticipantInfo]:
    participants = []
    for block in layout.participants:
        inn = ws.cell(row=8, column=block.start_col).value
        address = ws.cell(row=9, column=block.start_col).value
        accreditation = ws.cell(row=11, column=block.start_col).value
        participants.append(ParticipantInfo(
            id=block.participant_id,
            name=block.name,
            inn=str(inn).strip() if inn is not None else None,
            address=str(address).strip() if address is not None else None,
            accreditation_status=str(accreditation).strip() if accreditation else None,
        ))
    return participants


def _read_raw_rows(ws, last_row: int) -> list[RawRow]:
    rows = []
    for r in range(14, last_row + 1):
        b = ws.cell(row=r, column=2).value
        c = ws.cell(row=r, column=3).value
        d = ws.cell(row=r, column=4).value
        rows.append(RawRow(row=r, b=b, c=c, d=d))
    return rows


def _fill_position_data(ws, position: Position, layout: Layout, warnings: list[Warning]) -> None:
    row = position.row
    position.unit = ws.cell(row=row, column=7).value
    position.customer_quantity = _num(ws.cell(row=row, column=8).value, row, 8, warnings)
    comment = ws.cell(row=row, column=6).value
    position.customer_comment = str(comment).strip() if comment else None

    baseline_unit = _read_breakdown(ws, row, layout.baseline.unit_price_cols, warnings)
    baseline_total = _read_breakdown(ws, row, layout.baseline.total_cols, warnings)
    if any(v not in (None, 0) for v in (*baseline_unit.__dict__.values(), *baseline_total.__dict__.values())):
        position.baseline = BaselinePrice(unit_price=baseline_unit, total=baseline_total)

    for block in layout.participants:
        qty = _num(ws.cell(row=row, column=block.qty_col).value, row, block.qty_col, warnings)
        unit_price = _read_breakdown(ws, row, block.unit_price_cols, warnings)
        participant_total = _read_breakdown(ws, row, block.participant_total_cols, warnings)
        comment_raw = ws.cell(row=row, column=block.participant_comment_col).value
        p_comment = str(comment_raw).strip() if comment_raw else None

        status, refs = classify_status(unit_price.total, p_comment)

        customer_total = None
        if position.customer_quantity is not None and unit_price.total is not None:
            customer_total = position.customer_quantity * unit_price.total

        volume_mismatch = (
            qty is not None
            and position.customer_quantity is not None
            and abs(qty - position.customer_quantity) > 1e-6
        )
        if volume_mismatch:
            warnings.append(Warning(
                type=WarningType.VOLUME_MISMATCH, row=row, participant_id=block.participant_id,
                message=f"У {block.name} объём {qty} отличается от объёма заказчика {position.customer_quantity}",
            ))
        if status == PriceStatus.INCLUDED_ELSEWHERE:
            warnings.append(Warning(
                type=WarningType.INCLUDED_ELSEWHERE, row=row, participant_id=block.participant_id,
                message=f"{block.name}: {p_comment}",
            ))
        if status == PriceStatus.PLACEHOLDER:
            warnings.append(Warning(
                type=WarningType.PLACEHOLDER, row=row, participant_id=block.participant_id,
                message=f"{block.name}: цена-заглушка {unit_price.total}",
            ))

        pct = None
        if block.pct_col is not None:
            pct = _num(ws.cell(row=row, column=block.pct_col).value, row, block.pct_col, warnings)
        expected_cost = _num(
            ws.cell(row=row, column=block.expected_cost_col).value, row, block.expected_cost_col, warnings
        )

        position.participant_prices[block.participant_id] = PositionPrice(
            status=status,
            quantity_offered=qty,
            unit_price=unit_price,
            total_for_participant_volume=participant_total,
            total_for_customer_volume=customer_total,
            participant_comment=p_comment,
            mrg_comment=None,
            pct_of_estimate=pct,
            expected_cost=expected_cost,
            included_in_positions=refs,
            volume_mismatch=volume_mismatch,
        )


def _prune_empty_sections(sections: list[Section]) -> list[Section]:
    kept = []
    for section in sections:
        section.children = _prune_empty_sections(section.children)
        has_price = any(
            price.status != PriceStatus.ZERO
            for position in section.iter_positions()
            for price in position.participant_prices.values()
        )
        if has_price or section.children:
            kept.append(section)
    return kept


def _flag_lump_sums(sections: list[Section], participant_ids: list[str], warnings: list[Warning]) -> None:
    for section in sections:
        positions = list(section.iter_positions())
        if len(positions) < 2:
            continue
        for pid in participant_ids:
            priced_rows = [
                p.row for p in positions
                if p.participant_prices[pid].status == PriceStatus.PRICED
            ]
            if len(priced_rows) == 1:
                only_row = priced_rows[0]
                only_position = next(p for p in positions if p.row == only_row)
                old_price = only_position.participant_prices[pid]
                only_position.participant_prices[pid] = PositionPrice(
                    **{**old_price.__dict__, "status": PriceStatus.LUMP_SUM}
                )
                warnings.append(Warning(
                    type=WarningType.LUMP_SUM, row=only_row, participant_id=pid,
                    message=f"Участник дал весь раздел {section.number} одной строкой (п/п {only_row})",
                ))


def parse_tender(path: str | Path) -> Tender:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb[SHEET_NAME]

    layout = discover_layout(ws)
    warnings: list[Warning] = []

    subject, object_name, address = _read_header(ws)
    participants = _read_participants(ws, layout)

    raw_rows = _read_raw_rows(ws, ws.max_row)
    sections = build_section_tree(raw_rows)

    for section in sections:
        for position in section.iter_positions():
            _fill_position_data(ws, position, layout, warnings)

    sections = _prune_empty_sections(sections)
    _flag_lump_sums(sections, [p.id for p in participants], warnings)

    has_baseline = any(
        position.baseline is not None
        for section in sections
        for position in section.iter_positions()
    )
    mode = ComparisonMode.BASELINE if has_baseline else ComparisonMode.PEER

    return Tender(
        subject=subject, object_name=object_name, address=address,
        participants=participants, sections=sections, mode=mode, warnings=warnings,
    )
